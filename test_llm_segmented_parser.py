"""
Тестовый скрипт для LLM-Segmented Parser v1.0

Демонстрирует новую архитектуру:
1. LLM-сегментация на смысловые блоки
2. Детерминированный подбор правил по типу блока (вместо FAISS)
3. Поддержка PDF и DOCX
"""

import os
import json
from pathlib import Path

# Настраиваем логирование
os.environ.setdefault('LOG_TO_FILE', 'false')
os.environ.setdefault('LOG_LEVEL', 'INFO')

from src.llm_segmented_parser import LLMSegmentedParser, DeterministicRetriever, RULES_DB


def test_deterministic_retriever():
    """Тест детерминированного ретривера."""
    print("\n" + "=" * 70)
    print("ТЕСТ 1: Детерминированный ретривер правил")
    print("=" * 70)
    
    retriever = DeterministicRetriever()
    
    # Тестируем для разных типов блоков
    test_cases = [
        ("section_header", "Заголовок раздела"),
        ("table_block", "Таблица с данными"),
        ("text", "Основной текст документа"),
        ("title_page", "Титульный лист"),
        ("toc", "Содержание"),
    ]
    
    for chunk_type, description in test_cases:
        print(f"\n📋 Тип блока: {chunk_type.upper()} ({description})")
        rules = retriever.search("", chunk_type, "", top_k=5)
        print(f"   Найдено правил: {len(rules)}")
        for i, rule in enumerate(rules[:3], 1):
            print(f"   {i}. {rule['id']}: {rule['text'][:60]}...")


def test_pdf_parsing():
    """Тест парсинга PDF с LLM-сегментацией."""
    print("\n" + "=" * 70)
    print("ТЕСТ 2: Парсинг PDF с LLM-сегментацией")
    print("=" * 70)
    
    # Проверяем наличие тестового PDF
    test_pdf = Path("/workspace/src/МФАС.563563.003 ПМ.pdf")
    
    if not test_pdf.exists():
        print(f"⚠️  Тестовый PDF не найден: {test_pdf}")
        print("   Пропускаем тест PDF")
        return
    
    # Получаем API ключ
    api_key = os.getenv("OPENROUTER_API_KEY")
    if not api_key:
        print("⚠️  API ключ OPENROUTER_API_KEY не установлен")
        print("   Парсер будет работать в режиме fallback (без LLM)")
    
    parser = LLMSegmentedParser(api_key=api_key)
    
    print(f"\n📖 Обработка файла: {test_pdf.name}")
    chunks = parser.parse_pdf(str(test_pdf))
    
    print(f"\n✅ Извлечено чанков: {len(chunks)}")
    
    # Статистика по типам
    from collections import Counter
    type_stats = Counter(c.chunk_type for c in chunks)
    print(f"\n📊 Распределение типов блоков:")
    for block_type, count in type_stats.most_common():
        print(f"   {block_type.upper()}: {count}")
    
    # Примеры чанков для каждого типа
    print(f"\n📝 Примеры чанков:")
    shown_types = set()
    for chunk in chunks[:10]:
        if chunk.chunk_type not in shown_types:
            shown_types.add(chunk.chunk_type)
            print(f"\n   [{chunk.chunk_type.upper()}] Стр. {chunk.location.get('page', '?')}")
            print(f"   Текст: {chunk.text[:100]}...")
            print(f"   Правил: {chunk.metadata.get('rules_count', 0)}")
    
    # Сохраняем результаты
    output_path = "/workspace/output_llm_segmented.json"
    parser.save_chunks(chunks, output_path)
    print(f"\n💾 Результаты сохранены в: {output_path}")


def test_docx_parsing():
    """Тест парсинга DOCX (демонстрация интерфейса)."""
    print("\n" + "=" * 70)
    print("ТЕСТ 3: Парсинг DOCX (демонстрация)")
    print("=" * 70)
    
    # Создаём тестовый DOCX файл
    try:
        from docx import Document
        
        doc = Document()
        doc.add_heading('Отчет по практике', 0)
        doc.add_heading('1 Введение', level=1)
        doc.add_paragraph('Это тестовый документ для проверки парсера DOCX.')
        doc.add_paragraph('Числа от 1 до 9 пишутся словами, а с единицами измерения — цифрами.')
        doc.add_heading('2 Основная часть', level=1)
        doc.add_paragraph('Таблица 1 — Пример данных')
        
        table = doc.add_table(rows=3, cols=3)
        table.cell(0, 0).text = 'Параметр'
        table.cell(0, 1).text = 'Значение'
        table.cell(0, 2).text = 'Ед. изм.'
        table.cell(1, 0).text = 'Длина'
        table.cell(1, 1).text = '100'
        table.cell(1, 2).text = 'мм'
        
        doc.add_paragraph('Рисунок 1 — Схема процесса')
        
        test_docx = Path("/workspace/test_document.docx")
        doc.save(str(test_docx))
        
        api_key = os.getenv("OPENROUTER_API_KEY")
        parser = LLMSegmentedParser(api_key=api_key)
        
        print(f"\n📖 Обработка файла: {test_docx.name}")
        chunks = parser.parse_docx(str(test_docx))
        
        print(f"\n✅ Извлечено чанков: {len(chunks)}")
        
        # Статистика
        from collections import Counter
        type_stats = Counter(c.chunk_type for c in chunks)
        print(f"\n📊 Распределение типов блоков:")
        for block_type, count in type_stats.most_common():
            print(f"   {block_type.upper()}: {count}")
        
        # Сохраняем
        output_path = "/workspace/output_docx_llm_segmented.json"
        parser.save_chunks(chunks, output_path)
        print(f"\n💾 Результаты сохранены в: {output_path}")
        
        # Удаляем тестовый файл
        test_docx.unlink()
        
    except ImportError:
        print("⚠️  Библиотека python-docx не установлена")
        print("   Установите: pip install python-docx")
    except Exception as e:
        print(f"⚠️  Ошибка при тестировании DOCX: {e}")


def test_rule_coverage():
    """Тест покрытия правилами различных типов блоков."""
    print("\n" + "=" * 70)
    print("ТЕСТ 4: Покрытие правилами типов блоков")
    print("=" * 70)
    
    from src.llm_segmented_parser import BLOCK_RULE_REGISTRY
    
    print(f"\n📋 Всего типов блоков: {len(BLOCK_RULE_REGISTRY)}")
    print(f"📚 Всего правил в базе: {len(RULES_DB)}")
    
    print(f"\n📊 Правила по типам блоков:")
    for block_type, rule_ids in sorted(BLOCK_RULE_REGISTRY.items()):
        print(f"\n   {block_type}:")
        for rule_id in rule_ids:
            rule = RULES_DB.get(rule_id, {})
            rule_text = rule.get('text', '???')[:50]
            print(f"      - {rule_id}: {rule_text}...")


def main():
    """Запуск всех тестов."""
    print("\n" + "╔" + "═" * 68 + "╗")
    print("║" + " " * 15 + "LLM-SEGMENTED PARSER v1.0 - ТЕСТЫ" + " " * 16 + "║")
    print("╚" + "═" * 68 + "╝")
    
    # Тест 1: Детерминированный ретривер
    test_deterministic_retriever()
    
    # Тест 2: Парсинг PDF
    test_pdf_parsing()
    
    # Тест 3: Парсинг DOCX
    test_docx_parsing()
    
    # Тест 4: Покрытие правилами
    test_rule_coverage()
    
    print("\n" + "=" * 70)
    print("✅ ВСЕ ТЕСТЫ ЗАВЕРШЕНЫ")
    print("=" * 70 + "\n")


if __name__ == "__main__":
    main()
